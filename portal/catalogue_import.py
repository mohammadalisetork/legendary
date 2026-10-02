"""Validated XLSX catalogue template, preview, confirmation, and export helpers."""
import hashlib
from itertools import islice
from io import BytesIO

from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill

from .models import CatalogueImportBatch, Category, Department, Service, User
from .policies import eligible_owners

HEADERS = ["action", "family_slug", "service_code", "name", "domain", "short_description", "full_description",
           "purpose", "scope", "deliverables", "required_inputs", "request_requirements", "process_information", "excluded",
           "service_role", "acceptance_criteria", "legacy_sla", "initial_response_days", "review_target_days",
           "delivery_min_days", "delivery_max_days", "maximum_duration_days", "supports_desired_date", "default_owner_username", "active"]
REQUIRED = ["family_slug", "service_code", "name", "domain", "full_description", "initial_response_days", "delivery_min_days", "delivery_max_days", "active"]


def template_bytes():
    wb = Workbook()
    ws = wb.active
    ws.title = "Services"
    ws.append(HEADERS)
    ws.append(["CREATE", "family-slug", "SVC-001", "Sample service", "General", "", "Replace with service description", "", "", "", "", "", "", "", "", "", "", 2, "", 5, 10, "", "TRUE", "", "TRUE"])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="174F46")
    guide = wb.create_sheet("Instructions")
    for row in [("Field", "Rule"), ("action", "CREATE for new service, UPDATE for exact service_code match, or SKIP"),
                ("family_slug", "Existing family slug within the selected department"),
                ("service_code", "Stable unique code; exact match only"),
                ("default_owner_username", "Optional active manager eligible in this department"),
                ("active", "TRUE or FALSE"), ("Durations", "Positive whole days; optional targets may be blank; delivery_max_days must be >= delivery_min_days")]:
        guide.append(row)
    guide.freeze_panes = "A2"
    guide.column_dimensions["A"].width = 30
    guide.column_dimensions["B"].width = 100
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


def _int(value, label, errors):
    try:
        number = int(value)
        if number < 1: raise ValueError
        return number
    except (TypeError, ValueError):
        errors.append(f"{label}: باید عدد صحیح مثبت باشد")
        return None


def _optional_int(value, label, errors):
    if value in (None, ""): return None
    return _int(value, label, errors)


def validate_workbook(upload, department):
    if upload.size > 5 * 1024 * 1024:
        raise ValidationError("حداکثر اندازهٔ فایل اکسل ۵ مگابایت است.")
    if not upload.name.lower().endswith(".xlsx"):
        raise ValidationError("فقط فایل XLSX پذیرفته می‌شود.")
    content = upload.read()
    upload.seek(0)
    if not content.startswith(b"PK\x03\x04"):
        raise ValidationError("ساختار فایل XLSX معتبر نیست.")
    try:
        wb = load_workbook(BytesIO(content), read_only=True, data_only=False)
        ws = wb["Services"] if "Services" in wb.sheetnames else wb.active
        rows = ws.iter_rows(values_only=False)
        header_cells = next(rows, None)
        headers = [str(c.value or "").strip() for c in header_cells or []]
        if headers != HEADERS:
            raise ValidationError("ستون‌های فایل با قالب رسمی هم‌خوان نیستند؛ قالب را دوباره دریافت کنید.")
        families = {x.slug: x for x in department.service_families.all()}
        owners = {x.username: x for x in eligible_owners(department)}
        data_rows=list(islice(rows,10001))
        if len(data_rows)>10000:raise ValidationError("حداکثر ۱۰٬۰۰۰ ردیف در هر بار ورود پذیرفته می‌شود.")
        candidate_codes={str(cells[2].value).strip() for cells in data_rows if len(cells)>2 and cells[2].value not in (None,"")}
        existing_by_code=Service.objects.filter(code__in=candidate_codes).select_related("category__department").in_bulk(field_name="code")
        seen = set()
        parsed = []
        for line, cells in enumerate(data_rows, start=2):
            values = {}
            formula = False
            for index, key in enumerate(HEADERS):
                cell = cells[index] if index < len(cells) else None
                if cell and cell.data_type == "f": formula = True
                values[key] = cell.value if cell else None
            if all(value in (None, "") for value in values.values()): continue
            errors = []
            if formula: errors.append("فرمول در فایل اکسل مجاز نیست")
            normalized = {k: str(v).strip() if v is not None else "" for k, v in values.items()}
            action = normalized["action"].upper()
            code = normalized["service_code"]
            if action not in {"CREATE", "UPDATE", "SKIP"}: errors.append("action: فقط CREATE، UPDATE یا SKIP")
            for field in REQUIRED:
                if normalized[field] == "": errors.append(f"{field}: مقدار الزامی است")
            if not code or len(code) > 20: errors.append("service_code: حداکثر ۲۰ نویسه")
            if code in seen: errors.append("service_code: کد تکراری در همین فایل")
            seen.add(code)
            family = families.get(normalized["family_slug"])
            if not family: errors.append("family_slug: خانواده در این اداره پیدا نشد")
            elif family.lifecycle_status != Category.Status.ACTIVE: errors.append("family_slug: خانواده غیرفعال یا بایگانی شده است")
            owner = owners.get(normalized["default_owner_username"]) if normalized["default_owner_username"] else None
            if normalized["default_owner_username"] and not owner: errors.append("default_owner_username: مسئول فعال و مجاز این اداره پیدا نشد")
            initial = _int(values["initial_response_days"], "initial_response_days", errors)
            review = _optional_int(values["review_target_days"], "review_target_days", errors)
            minimum = _int(values["delivery_min_days"], "delivery_min_days", errors)
            maximum = _int(values["delivery_max_days"], "delivery_max_days", errors)
            max_duration = _optional_int(values["maximum_duration_days"], "maximum_duration_days", errors)
            active_value = normalized["active"].upper()
            if active_value not in {"TRUE", "FALSE", "YES", "NO", "1", "0"}: errors.append("active: باید TRUE یا FALSE باشد")
            desired_value = normalized["supports_desired_date"].upper()
            if desired_value not in {"", "TRUE", "FALSE", "YES", "NO", "1", "0"}: errors.append("supports_desired_date: باید TRUE یا FALSE باشد")
            existing = existing_by_code.get(code)
            if existing and existing.department.pk != department.pk: errors.append("service_code: این کد متعلق به ادارهٔ دیگری است")
            if action == "CREATE" and existing: errors.append("service_code: کد موجود است؛ برای ویرایش UPDATE را مشخص کنید")
            if action == "UPDATE" and not existing: errors.append("service_code: برای UPDATE باید خدمت موجود باشد")
            if action == "UPDATE" and existing and existing.lifecycle_status==Service.Status.ARCHIVED:errors.append("service_code: خدمت بایگانی‌شده قابل ویرایش یا فعال‌سازی نیست")
            if action == "SKIP" and existing and existing.department.pk != department.pk: errors.append("service_code: خدمت خارج از محدوده اداره است")
            if minimum and maximum and maximum < minimum: errors.append("delivery_max_days: باید برابر یا بیشتر از delivery_min_days باشد")
            record = {"row": line, "action": action, "service_code": code, "family_id": family.pk if family else None,
                      "name": normalized["name"], "domain": normalized["domain"], "short_description":normalized["short_description"], "full_description": normalized["full_description"],
                      "purpose":normalized["purpose"], "scope":normalized["scope"], "deliverables":normalized["deliverables"], "required_inputs":normalized["required_inputs"],
                      "request_requirements":normalized["request_requirements"], "process_information":normalized["process_information"], "excluded":normalized["excluded"],
                      "service_role":normalized["service_role"], "acceptance_criteria":normalized["acceptance_criteria"], "legacy_sla":normalized["legacy_sla"],
                      "initial_response_days": initial, "review_target_days":review, "delivery_min_days": minimum, "delivery_max_days": maximum,
                      "maximum_duration_days":max_duration, "supports_desired_date":desired_value not in {"FALSE","NO","0"},
                      "default_owner_id": owner.pk if owner else None, "active": active_value in {"TRUE", "YES", "1"},
                      "result": "ERROR" if errors else (action if action != "SKIP" else "SKIP"), "errors": errors}
            parsed.append(record)
        if len(parsed) > 10000: raise ValidationError("حداکثر ۱۰٬۰۰۰ ردیف در هر بار ورود پذیرفته می‌شود.")
        if not parsed: raise ValidationError("فایل هیچ ردیف خدمتی ندارد.")
        wb.close()
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError("خواندن فایل اکسل ممکن نشد؛ فایل را با قالب XLSX رسمی بازبینی کنید.") from exc
    return parsed, hashlib.sha256(content).hexdigest(), len(content)


def confirm_batch(batch, actor):
    created = updated = skipped = 0
    with transaction.atomic():
        batch=CatalogueImportBatch.objects.select_for_update().get(pk=batch.pk)
        if batch.status != CatalogueImportBatch.Status.PREVIEWED:raise ValidationError("این پیش‌نمایش قبلاً تعیین تکلیف شده است.")
        rows = batch.validation_result.get("rows", [])
        if any(row["errors"] for row in rows):raise ValidationError("تا زمان رفع همهٔ خطاها امکان تأیید ورود وجود ندارد.")
        department=Department.objects.select_for_update().get(pk=batch.department_id)
        if department.status==Department.Status.ARCHIVED:raise ValidationError("ورود به ادارهٔ بایگانی‌شده مجاز نیست.")
        for row in rows:
            if row["action"] == "SKIP": skipped += 1; continue
            family = Category.objects.select_for_update().get(pk=row["family_id"], department=department)
            if family.lifecycle_status!=Category.Status.ACTIVE:raise ValidationError(f"ردیف {row['row']}: خانواده در پیش‌نمایش غیرفعال شده است.")
            owner_id=row["default_owner_id"]
            if owner_id:
                owner=User.objects.get(pk=owner_id,is_active=True)
                from .policies import can_be_default_owner
                if not can_be_default_owner(owner,department):raise ValidationError(f"ردیف {row['row']}: مسئول دیگر مجاز این اداره نیست.")
            defaults = {"category": family, "name": row["name"], "domain": row["domain"],
                        "short_description":row["short_description"], "full_description": row["full_description"], "purpose":row["purpose"],"scope":row["scope"],
                        "deliverables":row["deliverables"],"required_inputs":row["required_inputs"],"request_requirements":row["request_requirements"],
                        "process_information":row["process_information"],"excluded":row["excluded"],"service_role":row["service_role"],
                        "acceptance_criteria":row["acceptance_criteria"],"legacy_sla":row["legacy_sla"],
                        "initial_response_days": row["initial_response_days"], "review_target_days":row["review_target_days"],
                        "delivery_min_days": row["delivery_min_days"], "delivery_max_days": row["delivery_max_days"],
                        "maximum_duration_days":row["maximum_duration_days"],"supports_desired_date":row["supports_desired_date"],
                        "default_owner_id": row["default_owner_id"], "active": row["active"],
                        "lifecycle_status": Service.Status.ACTIVE if row["active"] else Service.Status.DISABLED}
            if row["action"] == "CREATE":
                if Service.objects.filter(code=row["service_code"]).exists():raise ValidationError(f"ردیف {row['row']}: کد خدمت از زمان پیش‌نمایش ایجاد شده است.")
                obj = Service(code=row["service_code"], **defaults)
                obj.save(); created += 1
            else:
                obj = Service.objects.select_for_update().get(code=row["service_code"], category__department=batch.department)
                if obj.lifecycle_status==Service.Status.ARCHIVED:raise ValidationError(f"ردیف {row['row']}: خدمت از زمان پیش‌نمایش بایگانی شده است.")
                for key, value in defaults.items(): setattr(obj, key, value)
                obj.save(); updated += 1
        batch.status = CatalogueImportBatch.Status.CONFIRMED
        batch.confirmed_by = actor; batch.confirmed_at = timezone.now()
        batch.created_count = created; batch.updated_count = updated; batch.skipped_count = skipped
        batch.save(update_fields=["status", "confirmed_by", "confirmed_at", "created_count", "updated_count", "skipped_count"])
    return created, updated, skipped


def export_bytes(services):
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Services")
    ws.append(HEADERS)
    for s in services.select_related("category", "default_owner").iterator(chunk_size=500):
        values=["UPDATE", s.category.slug, s.code, s.name, s.domain, s.short_description,s.full_description,
                s.purpose,s.scope,s.deliverables,s.required_inputs,s.request_requirements,s.process_information,s.excluded,
                s.service_role,s.acceptance_criteria,s.legacy_sla,s.initial_response_days,s.review_target_days,
                s.delivery_min_days,s.delivery_max_days,s.maximum_duration_days,"TRUE" if s.supports_desired_date else "FALSE",
                s.default_owner.username if s.default_owner_id else "", "TRUE" if s.active else "FALSE"]
        cells=[]
        for value in values:
            cell=WriteOnlyCell(ws,value=value)
            if isinstance(value,str):cell.data_type="s"
            cells.append(cell)
        ws.append(cells)
    stream = BytesIO(); wb.save(stream); return stream.getvalue()
