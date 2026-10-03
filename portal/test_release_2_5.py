"""Deterministic Release 2.5 facts, scope, exports, and PDF regression tests."""
from datetime import date,datetime,time,timedelta
from io import BytesIO

from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from .models import AllocationPeriod, ApprovalCase, ApprovalStep, Category, CreditAllocation, CreditReservation, Department, PriorityPolicy, Program, Project, Request, RequestHistory, RoleAssignment, Service, ServiceFormField, User, WorkingCalendar
from .reporting import aging, analytics_scope, approval_metrics, completed_during, credit_metrics, facts_for, group, matrix, metrics, window_from_params


class ManagementReportingTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin=User.objects.create_user(username='bi-admin',role=User.Role.ADMIN,must_change_password=False)
        cls.executive=User.objects.create_user(username='bi-executive',must_change_password=False)
        cls.supervisor=User.objects.create_user(username='bi-supervisor',must_change_password=False)
        cls.lead=User.objects.create_user(username='bi-lead',must_change_password=False)
        cls.manager=User.objects.create_user(username='bi-manager',must_change_password=False)
        cls.program_manager=User.objects.create_user(username='bi-program',must_change_password=False)
        cls.project_manager=User.objects.create_user(username='bi-project',must_change_password=False)
        cls.requester=User.objects.create_user(username='bi-requester',must_change_password=False)
        cls.department=Department.objects.create(code='bi-dept-a',name='Department A',status=Department.Status.PUBLISHED)
        cls.foreign_department=Department.objects.create(code='bi-dept-b',name='Department B',status=Department.Status.PUBLISHED)
        cls.family=Category.objects.create(name='Historical Family',slug='bi-family-a',department=cls.department)
        cls.foreign_family=Category.objects.create(name='Foreign Family',slug='bi-family-b',department=cls.foreign_department)
        cls.service=Service.objects.create(code='BI-A',name='Historical Service',category=cls.family,domain='BI',full_description='Service')
        cls.foreign_service=Service.objects.create(code='BI-B',name='Foreign Service',category=cls.foreign_family,domain='BI',full_description='Service')
        ServiceFormField.objects.create(service=cls.service,key='purpose',label='Historical Purpose',field_type='text',options=[])
        cls.program=Program.objects.create(code='bi-program-a',name='Historical Program');cls.program.status=Program.Status.ACTIVE;cls.program.save()
        cls.foreign_program=Program.objects.create(code='bi-program-b',name='Foreign Program');cls.foreign_program.status=Program.Status.ACTIVE;cls.foreign_program.save()
        cls.project=Project.objects.create(code='bi-project-a',name='Historical Project',program=cls.program);cls.project.status=Project.Status.ACTIVE;cls.project.save()
        cls.foreign_project=Project.objects.create(code='bi-project-b',name='Foreign Project',program=cls.foreign_program);cls.foreign_project.status=Project.Status.ACTIVE;cls.foreign_project.save()
        for user,role,scope,target in ((cls.executive,RoleAssignment.Role.EXECUTIVE_VIEWER,RoleAssignment.ScopeType.GLOBAL,{}),
            (cls.supervisor,RoleAssignment.Role.SUPERVISOR,RoleAssignment.ScopeType.GLOBAL,{}),
            (cls.lead,RoleAssignment.Role.DEPARTMENT_LEAD,RoleAssignment.ScopeType.DEPARTMENT,{'department':cls.department}),
            (cls.manager,RoleAssignment.Role.REQUEST_MANAGER,RoleAssignment.ScopeType.DEPARTMENT,{'department':cls.department}),
            (cls.program_manager,RoleAssignment.Role.PROGRAM_MANAGER,RoleAssignment.ScopeType.PROGRAM,{'program':cls.program}),
            (cls.project_manager,RoleAssignment.Role.PROJECT_MANAGER,RoleAssignment.ScopeType.PROJECT,{'project':cls.project,'program':cls.program})):
            RoleAssignment.objects.create(user=user,role=role,scope_type=scope,**target)
        WorkingCalendar.objects.create(name='BI calendar',active=True,weekend_days=[],workday_start=time(0,0),workday_end=time(23,59))

    def setUp(self):
        self.now=timezone.make_aware(datetime(2026,10,2,12))
        self.window=window_from_params({'period':'custom','start':'2026-09-01','end':'2026-10-02'})

    def make_request(self,status=Request.Status.COMPLETED,*,service=None,program=None,project=None,submitted=None,completed=None,first=None,due=None,delivery=None,priority=Request.Priority.NORMAL):
        submitted=submitted or self.now-timedelta(days=8)
        if status==Request.Status.COMPLETED:
            completed=completed or submitted+timedelta(days=2)
            if first is None:first=submitted+timedelta(hours=1)
        return Request.objects.create(requester=self.requester,requesting_unit='Unit',service=service or self.service,program=program,project_entity=project,title='Visible summary',status=status,priority=priority,
            submitted_at=submitted,completed_at=completed,first_response_at=first,expected_initial_response_at=due or submitted+timedelta(days=1),
            estimated_delivery_max=delivery or (submitted+timedelta(days=5)).date(),request_data={'purpose':'Confidential answer'})

    def test_exact_kpi_denominator_overdue_aging_and_percentiles(self):
        for _ in range(6):self.make_request(program=self.program,project=self.project)
        for _ in range(2):self.make_request(Request.Status.SUBMITTED,submitted=self.now-timedelta(hours=1),due=self.now+timedelta(hours=10))
        self.make_request(Request.Status.IN_PROGRESS,submitted=self.now-timedelta(days=8),due=self.now-timedelta(days=4),delivery=(self.now+timedelta(days=2)).date())
        self.make_request(Request.Status.CANCELLED)
        facts=facts_for(analytics_scope(self.admin),self.window,{},now=self.now);k=metrics(facts)
        self.assertEqual((k['total'],k['completed'],k['open'],k['overdue']),(10,6,3,1))
        self.assertEqual(k['completion_rate'],66.7);self.assertEqual(k['resolution_median'],172800)
        self.assertIsNotNone(k['resolution_p90']);self.assertEqual(sum(row['total'] for row in aging(facts)),3)
        self.assertEqual(k['first_sla_count'],7)

    def test_at_risk_threshold_and_pause_union(self):
        submitted=self.now-timedelta(hours=10)
        risk=self.make_request(Request.Status.SUBMITTED,submitted=submitted,due=self.now+timedelta(hours=2))
        paused=self.make_request(Request.Status.SUBMITTED,submitted=submitted,due=self.now-timedelta(hours=4))
        start=submitted+timedelta(hours=2)
        case=ApprovalCase.objects.create(request=paused,trigger='PRIORITY',status=ApprovalCase.Status.PENDING,requested_by=self.requester)
        ApprovalStep.objects.create(case=case,request=paused,sequence=1,target='PROGRAM_MANAGER',status=ApprovalStep.Status.PENDING,pauses_sla=True,started_at=start)
        event=RequestHistory.objects.create(request=paused,actor=self.manager,action='CLOCK_PAUSED',to_status=Request.Status.NEED_INFO)
        RequestHistory.objects.filter(pk=event.pk).update(created_at=start)
        facts={f.request.pk:f for f in facts_for(analytics_scope(self.admin),self.window,{},now=self.now)}
        self.assertTrue(facts[risk.pk].at_risk);self.assertFalse(facts[risk.pk].overdue)
        self.assertFalse(facts[paused.pk].overdue);self.assertGreater(facts[paused.pk].program_wait,0)

    def test_historical_snapshots_and_legacy_unknown(self):
        old=self.make_request(program=self.program,project=self.project)
        Request.objects.filter(pk=old.pk).update(program_name_snapshot='Historical Program',project_name_snapshot='Historical Project')
        legacy=self.make_request(Request.Status.SUBMITTED)
        self.service.name='Renamed service';self.service.save();self.family.name='Renamed family';self.family.save()
        self.program.name='Renamed program';self.program.save();self.project.name='Renamed project';self.project.save()
        facts=facts_for(analytics_scope(self.admin),self.window,{},now=self.now)
        historic=next(f for f in facts if f.request.pk==old.pk)
        self.assertEqual((historic.service,historic.family,historic.program,historic.project),('Historical Service','Historical Family','Historical Program','Historical Project'))
        self.assertEqual(next(f for f in facts if f.request.pk==legacy.pk).program,'قدیمی / نگاشت‌نشده')
        self.assertEqual(group(facts,'service')[0]['name'],'Historical Service')

    def test_jalali_previous_period_and_timezone_boundary(self):
        w=window_from_params({'period':'month'},today=date(2026,10,2))
        self.assertEqual((w.start,w.previous().start),(date(2026,9,23),date(2026,8,23)))
        self.assertEqual((w.previous().end-w.previous().start).days,(w.end-w.start).days)
        with timezone.override('Asia/Tehran'):
            instant=timezone.make_aware(datetime(2026,10,1,0,30));self.make_request(Request.Status.SUBMITTED,submitted=instant)
            self.assertEqual(len(facts_for(analytics_scope(self.admin),window_from_params({'period':'today'},today=date(2026,10,1)),{},now=self.now)),1)

    def test_scoped_filters_matrix_and_drilldown(self):
        local=self.make_request(program=self.program,project=self.project)
        foreign=self.make_request(service=self.foreign_service,program=self.foreign_program,project=self.foreign_project)
        for user in (self.lead,self.program_manager,self.project_manager):self.assertEqual(len(facts_for(analytics_scope(user),self.window,{},now=self.now)),1)
        self.assertEqual(len(matrix(facts_for(analytics_scope(self.admin),self.window,{},now=self.now))),2)
        self.client.force_login(self.lead)
        self.assertEqual(self.client.get(reverse('analytics_dashboard')+'?department=bi-dept-b').status_code,404)
        self.assertNotIn(foreign.public_id,self.client.get(reverse('analytics_requests')).content.decode())
        self.client.force_login(self.program_manager)
        self.assertEqual(self.client.get(reverse('analytics_dashboard')+f'?program={self.foreign_program.pk}').status_code,404)
        self.assertEqual(self.client.get(reverse('analytics_program',args=[self.foreign_program.pk])).status_code,404)
        self.assertEqual(self.client.get(reverse('analytics_program',args=[self.program.pk])).status_code,200)
        self.assertEqual(local.department_id,self.department.pk)

    def test_empty_authorized_dashboard_and_forced_export_links(self):
        self.client.force_login(self.lead)
        page=self.client.get(reverse('analytics_department',args=[self.department.code]))
        self.assertEqual(page.status_code,200);self.assertEqual(page.context['selected']['department'],self.department.code)
        self.assertEqual(page.context['chart_data']['sla'],[])
        self.assertIn('department=bi-dept-a',page.context['request_urls']['total'])
        self.assertEqual(self.client.get(reverse('analytics_department',args=[self.foreign_department.code])).status_code,404)
        self.client.force_login(self.program_manager)
        self.assertEqual(self.client.get(reverse('analytics_program',args=[self.program.pk])).context['selected']['program'],str(self.program.pk))
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(reverse('analytics_service',args=[self.service.pk])).context['selected']['service'],str(self.service.pk))

    def test_executive_safe_drill_supervisor_read_only_requester_denied(self):
        req=self.make_request();self.client.force_login(self.executive)
        self.assertEqual(self.client.get(reverse('analytics_dashboard')).status_code,200)
        page=self.client.get(reverse('analytics_requests')+'?period=custom&start=2026-09-01&end=2026-10-02')
        self.assertContains(page,req.public_id);self.assertNotContains(page,'Confidential answer')
        self.assertNotContains(page,reverse('request_detail',args=[req.pk]))
        self.assertEqual(self.client.get(reverse('request_detail',args=[req.pk])).status_code,404)
        self.client.force_login(self.supervisor)
        self.assertEqual(self.client.get(reverse('analytics_dashboard')).status_code,200)
        self.assertNotEqual(self.client.post(reverse('control_action',args=[req.pk]),{'kind':'note','body':'No'}).status_code,200)
        self.client.force_login(self.requester)
        for route in ('analytics_dashboard','analytics_excel','analytics_pdf'):
            self.assertEqual(self.client.get(reverse(route)).status_code,404)

    def test_excel_pdf_audit_formula_and_pdf_authorization(self):
        self.make_request();self.service.name='=1+1';self.service.save();self.make_request()
        self.client.force_login(self.admin)
        query='?period=custom&start=2026-09-01&end=2026-10-02'
        result=self.client.get(reverse('analytics_excel')+query);self.assertEqual(result.status_code,200)
        book=load_workbook(BytesIO(result.content));self.assertEqual(book['Requests'].max_row,3);self.assertEqual(book['Requests']['G3'].data_type,'s')
        standard=self.client.get(reverse('analytics_pdf')+query);self.assertEqual(standard.status_code,200);self.assertTrue(standard.content.startswith(b'%PDF-'))
        custom=self.client.get(reverse('analytics_pdf')+'?mode=custom&sections=summary&sections=appendix')
        self.assertEqual(custom.status_code,200);self.assertTrue(custom.content.startswith(b'%PDF-'))
        from .models import ActivityLog
        self.assertEqual(ActivityLog.objects.filter(action__in=['ANALYTICS_EXCEL_EXPORTED','EXECUTIVE_REPORT_GENERATED','CUSTOM_REPORT_GENERATED']).count(),3)
        self.client.force_login(self.lead)
        self.assertEqual(self.client.get(reverse('analytics_pdf')+query+'&department=bi-dept-b').status_code,404)

    def test_scoped_excel_program_filter_and_manager_visibility(self):
        local=self.make_request(program=self.program,project=self.project)
        foreign=self.make_request(service=self.foreign_service,program=self.foreign_program,project=self.foreign_project)
        query='?period=custom&start=2026-09-01&end=2026-10-02'
        self.client.force_login(self.lead)
        workbook=load_workbook(BytesIO(self.client.get(reverse('analytics_excel')+query).content))
        self.assertEqual((workbook['Requests'].max_row,workbook['Requests']['A2'].value),(2,local.public_id))
        self.client.force_login(self.executive)
        page=self.client.get(reverse('analytics_dashboard')+query+f'&program={self.program.pk}')
        self.assertEqual(page.context['kpis']['total'],1)
        self.assertNotIn(foreign.public_id,self.client.get(reverse('analytics_requests')+query+f'&program={self.program.pk}').content.decode())
        self.assertEqual(len(facts_for(analytics_scope(self.manager),self.window,{},now=self.now)),1)

    def test_approval_credit_and_batched_query(self):
        item=self.make_request(program=self.program,project=self.project,priority=Request.Priority.EMERGENCY)
        case=ApprovalCase.objects.create(request=item,trigger='PRIORITY',status=ApprovalCase.Status.APPROVED,requested_by=self.requester)
        ApprovalStep.objects.create(case=case,request=item,sequence=1,target='PROGRAM_MANAGER',status=ApprovalStep.Status.APPROVED,started_at=item.submitted_at,decided_at=item.submitted_at+timedelta(hours=2))
        period=AllocationPeriod.objects.create(name='Q4',kind=AllocationPeriod.Kind.QUARTER,starts_on=date(2026,9,1),ends_on=date(2026,12,31))
        priority=PriorityPolicy.objects.get(code=Request.Priority.EMERGENCY)
        wallet=CreditAllocation.objects.create(program=self.program,department=self.department,priority=priority,period=period,quantity=3)
        CreditReservation.objects.create(request=item,allocation=wallet,status=CreditReservation.Status.CONSUMED)
        facts=facts_for(analytics_scope(self.admin),self.window,{},now=self.now)
        self.assertEqual(approval_metrics(facts)['approved'],1)
        self.assertEqual((credit_metrics(self.admin,{},self.window)[0]['quantity'],credit_metrics(self.admin,{},self.window)[0]['remaining']),(3,2))
        self.assertEqual(credit_metrics(self.requester,{},self.window),[])
        with CaptureQueriesContext(connection) as queries:facts_for(analytics_scope(self.admin),self.window,{},now=self.now)
        self.assertLessEqual(len(queries),7)

    def test_credit_independence_by_department(self):
        period=AllocationPeriod.objects.create(name='FY',kind=AllocationPeriod.Kind.YEAR,starts_on=date(2026,1,1),ends_on=date(2026,12,31))
        priority=PriorityPolicy.objects.get(code=Request.Priority.VERY_URGENT)
        local=CreditAllocation.objects.create(program=self.program,department=self.department,priority=priority,period=period,quantity=1)
        CreditAllocation.objects.create(program=self.program,department=self.foreign_department,priority=priority,period=period,quantity=4)
        item=self.make_request(program=self.program,priority=Request.Priority.VERY_URGENT)
        CreditReservation.objects.create(request=item,allocation=local,status=CreditReservation.Status.CONSUMED)
        rows=credit_metrics(self.admin,{'program':self.program.pk},self.window)
        self.assertEqual({r['department']:(r['quantity'],r['remaining']) for r in rows},{self.department.name:(1,0),self.foreign_department.name:(4,4)})
        self.assertEqual(credit_metrics(self.lead,{},self.window)[0]['department'],self.department.name)

    def test_previous_period_event_throughput(self):
        self.make_request(submitted=self.now-timedelta(days=40),completed=self.now-timedelta(days=39))
        self.make_request(submitted=self.now-timedelta(days=1),completed=self.now)
        window=window_from_params({'period':'30d'},today=self.now.date())
        previous=metrics(facts_for(analytics_scope(self.admin),window.previous(),{},now=self.now))
        current=metrics(facts_for(analytics_scope(self.admin),window,{},now=self.now))
        self.assertEqual((previous['total'],current['total'],current['total']-previous['total']),(1,1,0))
        self.assertEqual(completed_during(analytics_scope(self.admin),window,{})[self.now.date()],1)

    def test_analytics_shell_local_charts_and_scope_preserving_reset(self):
        self.make_request(program=self.program,project=self.project)
        self.client.force_login(self.lead)
        page=self.client.get(reverse('analytics_department',args=[self.department.code]))
        self.assertEqual(page.status_code,200)
        self.assertEqual(page.context['clear_url'],reverse('analytics_department',args=[self.department.code]))
        self.assertEqual(page.context['chart_data']['matrix'],page.context['matrix_rows'])
        self.assertContains(page,'vendor/echarts-5.6.0.min.js')
        self.assertContains(page,'data-report-chart="matrix"')
        self.assertContains(page,'توزیع وضعیت فعلی')
        self.client.force_login(self.admin)
        executive_shell=self.client.get(reverse('analytics_dashboard'))
        self.assertEqual(executive_shell.content.decode().count('data-tour-title="گزارش و تحلیل"'),1)

    def test_capacity_dashboard_shows_utilization_without_counting_released_as_used(self):
        item=self.make_request(program=self.program,project=self.project,priority=Request.Priority.EMERGENCY)
        period=AllocationPeriod.objects.create(name='Capacity test',kind=AllocationPeriod.Kind.QUARTER,starts_on=date(2026,9,1),ends_on=date(2026,12,31))
        wallet=CreditAllocation.objects.create(program=self.program,department=self.department,priority=PriorityPolicy.objects.get(code=Request.Priority.EMERGENCY),period=period,quantity=4)
        CreditReservation.objects.create(request=item,allocation=wallet,status=CreditReservation.Status.RELEASED)
        self.client.force_login(self.admin)
        page=self.client.get(reverse('capacity_dashboard'))
        self.assertEqual(page.status_code,200)
        row=next(row for row in page.context['rows'] if row['wallet'].pk==wallet.pk)
        self.assertEqual((row['released'],row['utilization_pct'],row['remaining']),(1,0,4))
        self.assertContains(page,'آزادشده')

    def test_department_setup_checklist_keeps_existing_actions(self):
        self.client.force_login(self.lead)
        page=self.client.get(reverse('manage_department_detail',args=[self.department.pk]))
        self.assertEqual(page.status_code,200)
        for label in ('راهنمای راه‌اندازی','مدیر اداره و اعضای عملیاتی','خانواده، خدمت و فرم پویا','بازبینی پیش‌نمایش و انتشار'):
            self.assertContains(page,label)
        self.assertContains(page,reverse('manage_department_preview',args=[self.department.pk]))
