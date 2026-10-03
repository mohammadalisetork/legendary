"""Local-only UAT data setup and review path regression tests."""
from datetime import date
from io import BytesIO, StringIO

from django.core.management import call_command, CommandError
from django.test import TestCase, override_settings
from django.utils import timezone
from django.urls import reverse
from openpyxl import load_workbook

from .models import ApprovalCase, Category, CreditAllocation, CreditReservation, Department, Request, RoleAssignment, Service, User
from .policies import Action, can, has_department_role
from .reporting import analytics_scope, approval_metrics, credit_metrics, facts_for, window_from_params


class LocalUATSeedTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        with override_settings(APP_ENV="development", DEBUG=True):
            call_command("seed_demo", confirm_local_only=True, stdout=StringIO())

    def test_seed_is_idempotent_and_keeps_governance_examples(self):
        before = (User.objects.filter(username__startswith="uat-").count(),
                  Category.objects.filter(slug__startswith="uat-").count(),
                  Service.objects.filter(code__startswith="UAT-").count(),
                  Request.objects.filter(public_id__startswith="UAT-RQ-").count(),
                  ApprovalCase.objects.filter(request__public_id__startswith="UAT-RQ-").count(),
                  CreditReservation.objects.filter(request__public_id__startswith="UAT-RQ-").count())
        with override_settings(APP_ENV="development", DEBUG=True):
            call_command("seed_demo", confirm_local_only=True, stdout=StringIO())
        after = (User.objects.filter(username__startswith="uat-").count(),
                 Category.objects.filter(slug__startswith="uat-").count(),
                 Service.objects.filter(code__startswith="UAT-").count(),
                 Request.objects.filter(public_id__startswith="UAT-RQ-").count(),
                 ApprovalCase.objects.filter(request__public_id__startswith="UAT-RQ-").count(),
                 CreditReservation.objects.filter(request__public_id__startswith="UAT-RQ-").count())
        self.assertEqual(after, before)
        self.assertEqual(Department.objects.filter(status=Department.Status.PUBLISHED).count(), 6)
        self.assertEqual(before[:4], (9, 6, 6, 18))
        self.assertEqual(CreditAllocation.objects.filter(program__code__startswith="uat-program-").count(), 36)
        admin = User.objects.get(username="uat-admin")
        year = timezone.localdate().year
        fact_rows = facts_for(analytics_scope(admin), window_from_params({
            "period": "custom", "start": date(year, 1, 1).isoformat(),
            "end": date(year, 12, 31).isoformat()}), {})
        self.assertIn("UAT-RQ-001", [fact.request.public_id for fact in fact_rows if fact.at_risk])
        self.assertIn("UAT-RQ-002", [fact.request.public_id for fact in fact_rows if fact.overdue])
        self.assertEqual(set(CreditReservation.objects.filter(request__public_id__startswith="UAT-RQ-").values_list("status", flat=True)),
                         {CreditReservation.Status.RESERVED, CreditReservation.Status.CONSUMED, CreditReservation.Status.RELEASED})

    def test_demo_login_rbac_analytics_filter_and_exports(self):
        self.assertTrue(self.client.login(username="uat-admin", password="Local-UAT-Only-2026!"))
        for username in ("uat-executive", "uat-supervisor", "uat-department-lead",
                         "uat-request-manager", "uat-program-manager", "uat-project-manager",
                         "uat-requester", "uat-senior-approver"):
            self.client.logout()
            self.assertTrue(self.client.login(username=username, password="Local-UAT-Only-2026!"), username)
        self.client.logout()
        self.assertTrue(self.client.login(username="uat-admin", password="Local-UAT-Only-2026!"))
        admin = User.objects.get(username="uat-admin")
        requester = User.objects.get(username="uat-requester")
        self.assertTrue(can(admin, Action.USER_MANAGE))
        self.assertTrue(has_department_role(User.objects.get(username="uat-request-manager"),
                                            RoleAssignment.Role.REQUEST_MANAGER,
                                            Department.objects.get(code="technical-support")))
        self.assertEqual(RoleAssignment.objects.filter(user=requester, role=RoleAssignment.Role.REQUESTER,
                                                        scope_type=RoleAssignment.ScopeType.GLOBAL).count(), 1)
        response = self.client.get(reverse("analytics_dashboard"))
        self.assertEqual(response.status_code, 200)
        self.assertGreater(response.context["kpis"]["total"], 0)
        year = timezone.localdate().year
        report_facts = facts_for(analytics_scope(admin), window_from_params({
            "period": "custom", "start": date(year, 1, 1).isoformat(),
            "end": date(year, 12, 31).isoformat()}), {})
        self.assertGreaterEqual(approval_metrics(report_facts)["total"], 4)
        self.assertGreaterEqual(len(credit_metrics(admin, {}, window_from_params({
            "period": "custom", "start": date(year, 1, 1).isoformat(),
            "end": date(year, 12, 31).isoformat()}))), 1)
        filtered = self.client.get(reverse("analytics_dashboard") + "?status=IN_PROGRESS")
        self.assertEqual(filtered.status_code, 200)
        self.assertGreater(filtered.context["kpis"]["total"], 0)
        query = f"?period=custom&start={year}-01-01&end={year}-12-31"
        workbook_response = self.client.get(reverse("analytics_excel") + query)
        self.assertEqual(workbook_response.status_code, 200)
        workbook = load_workbook(BytesIO(workbook_response.content))
        self.assertGreater(workbook["Requests"].max_row, 2)
        pdf_response = self.client.get(reverse("analytics_pdf") + query)
        self.assertEqual(pdf_response.status_code, 200)
        self.assertTrue(pdf_response.content.startswith(b"%PDF-"))
        self.assertEqual(self.client.get("/health").status_code, 200)
        self.client.logout()
        self.assertTrue(self.client.login(username="uat-requester", password="Local-UAT-Only-2026!"))
        self.assertEqual(self.client.get(reverse("service_hub")).status_code, 200)
        self.assertEqual(self.client.get(reverse("catalog")).status_code, 200)
        service = Service.objects.get(code="UAT-02")
        self.assertEqual(self.client.get(reverse("service_detail", args=[service.pk])).status_code, 200)
        self.assertEqual(self.client.get(reverse("analytics_dashboard")).status_code, 404)
        self.assertGreaterEqual(Request.objects.filter(requester=requester).count(), 18)

    def test_seed_guard_requires_explicit_local_mode(self):
        with self.assertRaises(CommandError):
            call_command("seed_demo", stdout=StringIO())
        with override_settings(APP_ENV="production", DEBUG=False):
            with self.assertRaises(CommandError):
                call_command("seed_demo", confirm_local_only=True, stdout=StringIO())
