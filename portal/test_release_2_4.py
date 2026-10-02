from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.exceptions import ValidationError
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from .catalogue_import import HEADERS, template_bytes
from .models import (ActivityLog, CatalogueImportBatch, Category, Department,
                    Request, RoleAssignment, Service, ServiceFormField, User)


class CatalogueScalingTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin=User.objects.create_user(username="cat-admin",role=User.Role.ADMIN,must_change_password=False)
        cls.lead=User.objects.create_user(username="cat-lead",role=User.Role.USER,must_change_password=False)
        cls.outsider=User.objects.create_user(username="cat-outsider",role=User.Role.USER,must_change_password=False)
        cls.dept=Department.objects.create(code="cat-dept",name="Catalogue Dept",status=Department.Status.PUBLISHED)
        cls.other=Department.objects.create(code="cat-other",name="Other Dept",status=Department.Status.PUBLISHED)
        cls.family=Category.objects.create(name="Family",slug="family",department=cls.dept)
        cls.other_family=Category.objects.create(name="Other Family",slug="other-family",department=cls.other)
        cls.service=Service.objects.create(code="CAT-1",name="Service",category=cls.family,domain="Core",full_description="Description")
        RoleAssignment.objects.create(user=cls.lead,role=RoleAssignment.Role.DEPARTMENT_LEAD,scope_type=RoleAssignment.ScopeType.DEPARTMENT,department=cls.dept)

    def workbook_upload(self, rows, name="catalogue.xlsx"):
        wb=Workbook();ws=wb.active;ws.title="Services";ws.append(HEADERS)
        for row in rows:ws.append(row)
        stream=BytesIO();wb.save(stream)
        return SimpleUploadedFile(name,stream.getvalue(),content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def import_row(self,action,code="CAT-NEW",family="family",name="New service",minimum=3,maximum=7,active="TRUE"):
        return [action,family,code,name,"Core","Short description","Full description","Purpose","Scope","Deliverables","Inputs","Requirements","Process","Excluded","Role","Acceptance","Legacy SLA",2,"",minimum,maximum,"","TRUE","",active]

    def test_department_publish_requires_active_family_service_and_temporary_disable_is_non_requestable(self):
        draft=Department.objects.create(code="empty-dept",name="Empty")
        from .forms import DepartmentLifecycleForm
        form=DepartmentLifecycleForm({"status":Department.Status.PUBLISHED},department=draft,actor=self.admin)
        self.assertFalse(form.is_valid())
        self.dept.status=Department.Status.TEMPORARILY_DISABLED;self.dept.save(update_fields=["status"])
        self.assertFalse(self.dept.is_requestable)
        self.dept.status=Department.Status.PUBLISHED;self.dept.save(update_fields=["status"])
        self.assertTrue(self.dept.is_requestable)

    def test_family_and_service_lifecycle_soft_disable_without_deleting_history(self):
        item=Request.objects.create(requester=self.lead,requesting_unit="Unit",department=self.dept,service=self.service,title="Historical")
        self.service.lifecycle_status=Service.Status.ARCHIVED;self.service.active=False;self.service.save()
        self.family.lifecycle_status=Category.Status.ARCHIVED;self.family.active=False;self.family.save()
        self.assertTrue(Request.objects.filter(pk=item.pk).exists())
        self.assertFalse(self.service.is_requestable)
        self.client.force_login(self.lead)
        self.assertEqual(self.client.get(reverse("request_create",args=[self.service.pk])).status_code,404)

    def test_form_builder_scope_validation_ordering_and_historical_snapshot(self):
        field=ServiceFormField.objects.create(service=self.service,key="details",label="Old label",field_type="text",display_order=2)
        later=ServiceFormField.objects.create(service=self.service,key="later",label="Later",field_type="text",display_order=7)
        choice=ServiceFormField.objects.create(service=self.service,key="choice",label="Choice",field_type="radio",options=["A","B"],display_order=1)
        phone=ServiceFormField.objects.create(service=self.service,key="phone",label="Phone",field_type="phone",display_order=3)
        req=Request.objects.create(requester=self.lead,requesting_unit="Unit",department=self.dept,service=self.service,title="Snapshot",request_data={"details":"Value"})
        self.assertEqual(dict(req.brief_items)["Old label"],"Value")
        self.client.force_login(self.lead)
        response=self.client.post(reverse("service_field_edit",args=[self.service.pk,field.pk]),{"key":"details","label":"New label","field_type":"text","required":"on","placeholder":"","help_text":"","options":"[]","display_order":0,"active":"on"})
        self.assertEqual(response.status_code,302)
        req.refresh_from_db();self.assertEqual(dict(req.brief_items)["Old label"],"Value")
        dup=self.client.post(reverse("service_field_create",args=[self.service.pk]),{"key":"details","label":"Dup","field_type":"text","options":"[]"})
        self.assertEqual(dup.status_code,200)
        self.assertEqual(self.client.get(reverse("service_fields",args=[self.service.pk])).status_code,200)
        preview=self.client.get(reverse("service_form_preview",args=[self.service.pk]))
        self.assertEqual(preview.status_code,200)
        self.assertEqual(preview.context["form"]["data_choice"].field.widget.__class__.__name__,"RadioSelect")
        self.assertEqual(preview.context["form"]["data_phone"].field.widget.input_type,"tel")
        self.assertEqual(list(self.service.form_fields.values_list("key",flat=True)),["details","choice","phone","later"])

    def test_department_lead_cannot_modify_import_or_export_other_department(self):
        self.client.force_login(self.lead)
        foreign=Service.objects.create(code="CAT-2",name="Foreign",category=self.other_family,domain="Core",full_description="Other")
        self.assertEqual(self.client.get(reverse("service_fields",args=[foreign.pk])).status_code,403)
        self.assertEqual(self.client.get(reverse("catalogue_import",args=[self.other.pk])).status_code,403)
        self.assertEqual(self.client.get(reverse("catalogue_export")+"?department=cat-other").status_code,404)
        self.service.name="=1+1";self.service.save()
        exported=self.client.get(reverse("catalogue_export")+"?department=cat-dept")
        self.assertEqual(exported.status_code,200)
        book=load_workbook(BytesIO(exported.content),data_only=False)
        service_name_cell=book["Services"]["D2"]
        self.assertEqual(service_name_cell.value,"=1+1");self.assertNotEqual(service_name_cell.data_type,"f")

    def test_excel_template_preview_is_read_only_and_confirm_is_audited(self):
        self.client.force_login(self.admin)
        template=self.client.get(reverse("catalogue_import_template",args=[self.dept.pk]))
        self.assertEqual(template.status_code,200);self.assertTrue(template.content.startswith(b"PK"))
        upload=self.workbook_upload([self.import_row("CREATE")])
        preview=self.client.post(reverse("catalogue_import",args=[self.dept.pk]),{"file":upload})
        self.assertEqual(preview.status_code,200)
        self.assertFalse(Service.objects.filter(code="CAT-NEW").exists())
        batch_id=preview.context["preview"].pk
        result=self.client.post(reverse("catalogue_import",args=[self.dept.pk]),{"confirm_batch":batch_id})
        self.assertEqual(result.status_code,302)
        self.assertTrue(Service.objects.filter(code="CAT-NEW",category=self.family).exists())
        self.assertTrue(ActivityLog.objects.filter(action="CATALOGUE_IMPORT_CONFIRMED").exists())

    def test_excel_validation_rejects_duplicate_codes_bad_durations_and_unauthorized_import(self):
        self.client.force_login(self.lead)
        duplicate=self.import_row("CREATE")
        response=self.client.post(reverse("catalogue_import",args=[self.dept.pk]),{"file":self.workbook_upload([duplicate,duplicate])})
        self.assertEqual(response.status_code,200)
        self.assertFalse(Service.objects.filter(code="CAT-NEW").exists())
        self.assertGreater(response.context["preview"].validation_result["errors"],0)
        invalid=self.import_row("CREATE",code="CAT-BAD",minimum=10,maximum=2)
        response=self.client.post(reverse("catalogue_import",args=[self.dept.pk]),{"file":self.workbook_upload([invalid])})
        self.assertGreater(response.context["preview"].validation_result["errors"],0)
        self.client.force_login(self.outsider)
        self.assertEqual(self.client.get(reverse("catalogue_import",args=[self.dept.pk])).status_code,403)

    def test_malformed_oversized_and_invalid_owner_workbooks_are_rejected_without_service_writes(self):
        self.client.force_login(self.admin)
        malformed=SimpleUploadedFile("bad.xlsx",b"not a workbook",content_type="application/octet-stream")
        response=self.client.post(reverse("catalogue_import",args=[self.dept.pk]),{"file":malformed})
        self.assertContains(response,"فایل قابل پردازش نیست")
        oversized=SimpleUploadedFile("large.xlsx",b"x"*(5*1024*1024+1),content_type="application/octet-stream")
        response=self.client.post(reverse("catalogue_import",args=[self.dept.pk]),{"file":oversized})
        self.assertContains(response,"۵ مگابایت")
        self.assertEqual(CatalogueImportBatch.objects.count(),0)
        row=self.import_row("CREATE",code="CAT-OWNER-IMPORT")
        row[-2]="cat-outsider"
        response=self.client.post(reverse("catalogue_import",args=[self.dept.pk]),{"file":self.workbook_upload([row])})
        self.assertGreater(response.context["preview"].validation_result["errors"],0)
        self.assertFalse(Service.objects.filter(code="CAT-OWNER-IMPORT").exists())

    def test_import_update_skip_and_conflict_matching_are_explicit(self):
        self.client.force_login(self.admin)
        Service.objects.create(code="CAT-2",name="Unchanged",category=self.family,domain="Core",full_description="Description")
        upload=self.workbook_upload([self.import_row("UPDATE",code="CAT-1",name="Updated"),self.import_row("SKIP",code="CAT-2")])
        response=self.client.post(reverse("catalogue_import",args=[self.dept.pk]),{"file":upload})
        self.assertIsNotNone(response.context["preview"],str(list(response.context["messages"])) if "messages" in response.context else response.content.decode()[-1000:])
        self.assertEqual(response.context["preview"].validation_result["errors"],0)
        self.client.post(reverse("catalogue_import",args=[self.dept.pk]),{"confirm_batch":response.context["preview"].pk})
        self.service.refresh_from_db();self.assertEqual(self.service.name,"Updated")
        conflict=self.client.post(reverse("catalogue_import",args=[self.dept.pk]),{"file":self.workbook_upload([self.import_row("CREATE",code="CAT-1")])})
        self.assertGreater(conflict.context["preview"].validation_result["errors"],0)

    def test_invalid_owner_is_rejected_by_service_model(self):
        with self.assertRaises(ValidationError):
            Service.objects.create(code="CAT-OWNER",name="Bad owner",category=self.family,domain="Core",full_description="Description",default_owner=self.outsider)
